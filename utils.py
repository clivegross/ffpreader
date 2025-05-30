from openpyxl import load_workbook
import pandas as pd
from openpyxl.worksheet.table import Table, TableStyleInfo


def load_text(filename):
    """load text file, return as string"""
    with open(filename, "r") as file:
        text = file.read()
    return text


def filter_sections_start(sections, keyword):
    """read list of strings and only return those that contain a keyword at start of string"""
    filtered = []
    for section in sections:
        if section.startswith(keyword):
            filtered.append(section)
    return filtered


def filter_sections_end(sections, keyword):
    """read list of strings and only return those that contain a keyword at end of first line"""
    filtered = []
    for section in sections:
        if section.split("\n")[0].endswith(keyword):
            filtered.append(section)
    return filtered


def parse_tsv(text):
    """
    Read a string, skip first line and parse the rest as a tab separated value (tsv) list
    Example:
    [ M 140102 X 2
    22	IRD-ICG-L00-CSV-01 BOH CORRIDOR	x02	OPT	0	0	N	Y	Y	Y	Y	Y	N	Y	6	0	0	0	0	180	N	100	80	0	0	0	0	0	NA
    22	IRD-IT1-L00-SEL-02 ELEC CUPBOARD	x02	OPT	0	0	N	Y	Y	Y	Y	Y	N	Y	6	0	0	0	0	180	N	100	80	0	0	0	0	0	NA...
    ]
    Returns:
    skips the first header line and treats the remaining string as a tab separated values (tsv) format data and returns as a list of lists
    [
        ['22', 'IRD-ICG-L00-CSV-01 BOH CORRIDOR', 'x02', 'OPT', '0', '0', 'N', 'Y', 'Y', 'Y', 'Y', 'Y', 'N', 'Y', '6', '0', '0', '0', '0', '180', 'N', '100', '80', '0', '0', '0', '0', '0', 'NA'],
        ['22', 'IRD-IT1-L00-SEL-02 ELEC CUPBOARD', 'x02', 'OPT', '0', '0', 'N', 'Y', 'Y', 'Y', 'Y', 'Y', 'N', 'Y', '6', '0', '0', '0', '0', '180', 'N', '100', '80', '0', '0', '0', '0', '0', 'NA...']
    ]
    """
    lines = text.split("\n")[1:]
    table = [line.split("\t") for line in lines]
    return table


def parse_tsvs(sections):
    """function to parse_tsv a list of strings"""
    tables = [parse_tsv(section) for section in sections]
    return tables


def list_to_df(list_of_lists, col_names=None):
    """convert list of lists to pandas dataframe"""
    if col_names is None:
        df = pd.DataFrame(list_of_lists)
    else:
        df = pd.DataFrame(list_of_lists, columns=col_names)
    return df


def combine_named_dfs_from_list_of_dicts(merged, data_key="devices"):
    """
    Combine all dfs in the 'data_key' column of a list of dicts into a single df
    takes a list of dicts containing a 'data_key' key with a pd.DataFrame value and combines them into a single pd.DataFrameq
    Example:
    merged = [
        {'devices': pd.DataFrame(...), 'id': '110101', 'loop': 23, 'node': 11, 'raw': 'M 110101 X 1'},
        {'devices': pd.DataFrame(...), 'id': 'xxxx', 'loop': 24, 'node': 11, 'raw': 'M xxxx X 1'},
        ...
    ]
    Returns:
    pd.DataFrame(...)
    """
    devices_dfs = [loop[data_key] for loop in merged]
    combined = pd.concat(devices_dfs)
    return combined


def write_dfs_to_excel_and_format(
    data, filepath, sheet_name_key="description", data_key="data", as_table=False
):
    """
    Write structured dict or list of dicts of DataFrames to an Excel file and format as a table.

    This function expects the input `data` to be in one of the following formats:
    1. A dictionary where each key maps to a list of dictionaries, each containing a DataFrame and metadata.
    2. A dictionary where each key maps directly to a DataFrame.
    3. A list of dictionaries, where each dictionary contains a DataFrame and metadata.

    The function normalizes the input data into the following format for processing:
    [
        {
            sheet_name_key: 'zones',
            data_key: pandas.DataFrame(...)
        },
        ...
    ]

    Args:
        data (dict or list): The input data to process.
        filepath (str): The path to the Excel file to write.
        sheet_name_key (str, optional): The key in each dictionary to use as the sheet name.
            Defaults to 'description'.
        data_key (str, optional): The key in each dictionary that contains the DataFrame.
            Defaults to 'data'.
        as_table (bool, optional): Whether to format the data as a table in Excel. Defaults to False.

    Raises:
        ValueError: If the input data is not in a supported format.
    """
    # Normalize the input data into a list of dictionaries
    if isinstance(data, dict):
        flattened_data = []
        for key, value in data.items():
            if isinstance(value, list):  # If value is a list of dicts
                flattened_data.extend(value)
            elif isinstance(value, pd.DataFrame):  # If value is a DataFrame
                flattened_data.append({sheet_name_key: key, data_key: value})
            else:
                raise ValueError(
                    f"Unsupported value type for key '{key}': {type(value)}. Expected list or DataFrame."
                )
    elif isinstance(data, list):  # If data is already a list of dicts
        flattened_data = data
    else:
        raise ValueError(
            "Input data must be a dict of lists of dicts, a dict of DataFrames, or a list of dicts."
        )

    # Write the DataFrames to Excel
    writer = pd.ExcelWriter(filepath, engine="openpyxl")
    for sheet_data in flattened_data:
        sheet_name = sheet_data.get(sheet_name_key)
        df = sheet_data.get(data_key)

        if not isinstance(df, pd.DataFrame):
            raise ValueError(
                f"Expected a DataFrame for key '{data_key}', but got {type(df)}."
            )
        if not isinstance(sheet_name, str):
            raise ValueError(
                f"Expected a string for key '{sheet_name_key}', but got {type(sheet_name)}."
            )

        print(f"Writing DataFrame to Excel sheet '{sheet_name}' in '{filepath}'...")
        df.to_excel(writer, index=False, sheet_name=sheet_name)

    writer.close()

    # Optionally format the data as Excel tables
    if as_table:
        book = load_workbook(filepath)
        for sheet_data in flattened_data:
            sheet_name = sheet_data.get(sheet_name_key)
            sheet = book[sheet_name]

            print(f"Adding Excel table to sheet '{sheet_name}'...")
            dimensions = sheet.calculate_dimension()
            table = Table(displayName=sheet_name, ref=dimensions)
            style = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            table.tableStyleInfo = style
            sheet.add_table(table)

        book.save(filepath)

    print("Done writing to Excel.")
